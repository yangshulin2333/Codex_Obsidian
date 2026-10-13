from pathlib import Path
import hashlib
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
QA = Path(__file__).resolve().parent
SOURCE = Path('C:/Users/long001/Documents/ChatGPT/Travelbuddy/exports/海箱旅蛙_Demo对齐_20261012')
FILES = {
    'game': ('01_完整游戏策划案_Demo对齐版.docx', '01_完整游戏策划案_Demo自主出游版.docx'),
    'program': ('02_程序执行方案_Demo对齐版.docx', '02_程序执行方案_Demo自主出游版.docx'),
    'art': ('03_美术需求方案_Demo对齐版.docx', '03_美术需求方案_Demo自主出游版.docx'),
}

RULES = [
 ('职责', '玩家经营岛屿、购买并装载补给、升级、开箱和出售；伙伴自行选择出发时间、已解锁海域和短长航程。伙伴是通用角色概念，海蛙是第一位伙伴。'),
 ('正常准备', '条件满足后随机等待 60～180 秒；本次海域、航程和出发时间保存，在线离线都推进，重进不重抽或重置等待。普通海域及短长航程的选择权重在配置表维护，具体比例尚未冻结。'),
 ('休息与首航', '实际归家后休息 180 秒，再进入下一次 60～180 秒准备。首航准备 30 秒、航行 60 秒、收益倍率 1。测试准备窗口保留 15～30 秒，其余时间沿用测试倍率；正式与测试参数分离。'),
 ('玩家愿望', '每趟最多一次有效改地点建议：80% 采纳，20% 保留原决定。原本同一海域直接同意；未解锁海域说明原因，不消耗有效建议机会。结果保存且只影响下一趟，途中不改航线。'),
 ('手动暂缓', '出发前暂缓必定生效；暂停标记持久保存，只有玩家明确恢复才解除。取空行囊取消出发准备，但保留路线决定，重新装入后生成一次等待，避免反复取放刷路线。'),
 ('行囊', '仓库补给与行囊补给分开；购买只增加仓库。行囊固定最多 3 份，装入按钮默认转移 1 份，需玩家主动操作；也可装满或收回。伙伴在家休息、准备时可调整，在途和未靠岸时只读。每趟出发扣 1 份，在途行囊归零仍正常完成。'),
 ('容量与停止', '住宅箱位 Lv1 为 3、Lv2 为 6，行囊始终为 3。占用量为未开封箱加在途预约，每趟出发预约 1 位，归家原子转换为 1 只封箱；只有开箱释放箱位。无行囊补给、箱满、船施工时待命，条件恢复后按休息及等待时间再出发。'),
 ('离线多趟', '离线也可在家形成决定、出发、归家、休息并继续下一趟。沿原绝对时间线补算，受真实经过时间、行囊和空箱位限制；一次装载最多三趟，包括已经扣补给的在途航次。仓库库存不会自动补入行囊。'),
 ('归家入库', '归家自动将封箱放入箱区、特产按数量合并到背包、明信片记入收藏。商品和木材开箱后入库；特产及商品发现、实际卡片 ID 与已开箱品质分别解锁。查看记录只标记已读，不发奖、不释放箱位、不阻止下一趟。'),
 ('旅行记录', '保存最近 10 趟详细记录，另保存尚未查看的收获汇总。记录包括航次、海域、时间、明信片、特产和对应箱 ID。重进显示新完成趟数、收获、行囊剩余、箱区占用、待命原因及当前在途倒计时；旧记录清理不影响资产、未读汇总和防重复结算。'),
 ('品质保密', '封箱只显示中性箱体及海域；未开箱不展示品质、商品、木材内容或估值。日志和离线小结同样遵守。每趟唯一航次 ID，补给、预约、入库和结算标记保证只提交一次。'),
 ('岛屿与表现', '四人共享一座主岛，每人独立住宅、船位和箱区，共享港口和商店，拍卖区为白模。岛上交互为 3D；伙伴离港后玩家仍可在岛上活动，主动打开可关闭的 2D 旅行观察窗。远方不制作真实 3D 关卡，明信片使用 2D 插画。'),
]

PARAMS = [
 ('正常准备等待', '60～180 秒', '随机一次并保存截止时间'),
 ('实际归家休息', '180 秒', '与查看收获无关'),
 ('有效地点建议', '80% 采纳', '20% 保留原决定 每趟一次'),
 ('手动暂缓', '必定生效', '玩家恢复前持久暂停'),
 ('行囊上限', '3 份', '每趟出发消耗 1 份'),
 ('箱区容量', 'Lv1 3 位 Lv2 6 位', '包含在途预约'),
 ('详细旅行记录', '最近 10 趟', '未读汇总单独保留'),
 ('首航', '准备 30 秒 航行 60 秒', '收益倍率固定 1'),
 ('测试准备窗口', '15～30 秒', '单独配置 避免随倍率缩到不可干涉'),
]

BASELINE = [
 ('初始资产', '30 金币、仓库补给 3 份、行囊 0 份、木材 4；住宅和船均 Lv1。购买补给 5 金币/份。'),
 ('住宅升级', '60 金币 + 8 木材，施工 180 秒；箱位 3→6，换较大木屋。可与旅行并行。'),
 ('船升级', '40 金币 + 6 木材，施工 120 秒；木船换帆船并解锁近海，施工期间禁止新出发，在途不能施工。'),
 ('升级余量', '扣费后保留至少 1 份仓库或行囊补给，或至少 5 金币。'),
 ('正式航程', '浅滩短航 600 秒/收益倍率 1，长航 1800 秒/倍率 3；近海短航 1800 秒/倍率 1，长航 3600 秒/倍率 2。伙伴从合法配置中选择。'),
 ('奖励数量', '每趟 1 箱、1 张明信片；特产数量为收益倍率。开箱抽取 3×倍率 件商品，固定 2 木材。首航倍率 1。'),
 ('品质权重', '普通、生活、精良三品质；浅滩 90/10/0%，近海 20/20/60%。品质、奖励快照留在服务端，船升级不直接改变海域权重。'),
 ('内容与出售', '12 件箱内商品、3 件特产、10 张明信片。固定售价沿用既定 15 个条目；出售后保留图鉴。'),
 ('应急补给', '金币不足 5，仓库及行囊均无补给，背包无可售资产、无封箱且无在途时，可领 1 份到仓库；仍需装入行囊。不能反复领或出售补给。'),
 ('后续范围', 'P0 不做菜地、交易拍卖、动态市场、付费、真实远海驾驶、复杂性格/服从度和途中召回。P1/P2 的独立规划保留，不作为本轮制作任务。'),
]

STEPS = [
 ('M0 统一目标', '同步自主出游、行囊、暂停、收获记录与配置口径', '正文、状态、接口、UI 和验收一致'),
 ('M1 存档扩展', '仓库/行囊分开，保存决定、暂停、随机序列、箱位预约、航次提交标记和未读汇总', '恢复旧存档不覆盖资产，重进不重抽'),
 ('M2 行囊交互', '购买、装入、装满、收回与在家校验', '每次转移数量守恒，重复请求不重复扣增'),
 ('M3 自主决定', '准备计时、合法海域/航程、一次愿望、暂缓和恢复', '80/20 只抽一次，离线准备仍推进'),
 ('M4 归家与离线', '绝对时间线、多趟推进、自动入库、休息及容量转换', '最多三趟，已在途归家不被无补给或箱满拦截'),
 ('M5 收获与经营', '未读小结、最近10趟记录、开箱出售、升级和图鉴回归', '查看不发奖，日志裁剪不影响资产'),
 ('M6 世界与界面', '四人岛屿、伙伴意图/待命提示、行囊和收获面板、2D 观察窗', 'PC/手机可完成准备与干涉，正式资产可后续替换'),
 ('M7 验收', '准备前/中离线、在途、多趟、满箱、耗尽、暂缓、重进和存档失败', '通过新规则验收后再进入后续版本'),
]

def paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs

def replace(doc, old, new):
    count = 0
    for p in paragraphs(doc):
        if old in p.text:
            p.text = p.text.replace(old, new)
            count += 1
    return count

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    for c, value in zip(table.rows[0].cells, headers):
        c.text = value
    for values in rows:
        cells = table.add_row().cells
        for c, value in zip(cells, values):
            c.text = value
    if widths:
        table.autofit = False
        for gc, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
            gc.set(qn('w:w'), str(round(width / 2.54 * 1440)))
        for row in table.rows:
            for c, width in zip(row.cells, widths):
                c.width = Cm(width)
    return table

def front(doc, title):
    body = doc.element.body
    # All three baseline files share the first nine blocks; replace their old rule table and gates.
    for block in list(body)[:9]:
        body.remove(block)
    tmp = Document()
    tmp.add_paragraph(title, 'Title')
    tmp.add_paragraph('Demo 自主出游版 V0.2 2026年10月13日')
    tmp.add_paragraph('本版确立玩家经营岛屿与准备补给、伙伴自主出游的第一版功能基准。伙伴可在线和离线决定并完成旅行；玩家通过行囊预算、一次地点愿望与暂缓参与。程序接入状态：本次更新方案文档，新规则尚待开发和验收。')
    tmp.add_heading('第一版统一职责与规则', level=1)
    # Use prose for policies; do not force long narrative blocks into table cells.
    for label, value in RULES:
        p = tmp.add_paragraph()
        p.add_run(label + '。').bold = True
        p.add_run(value)
    tmp.add_heading('第一版参数基准', level=1)
    add_table(tmp, ['参数', '数值', '说明'], PARAMS, [4.0, 4.5, 7.9])
    tmp.add_heading('已确认的经营与航程配置', level=1)
    add_table(tmp, ['项目', '规则'], BASELINE, [3.1, 13.3])
    tmp.add_heading('新规则开发顺序与交付门槛', level=1)
    tmp.add_paragraph('先扩展现有业务存档和规则，再接入行囊、自主决定、自动归家与离线多趟；最后改界面并回归经营和联机。通用框架继续复用，正式美术按功能验证结果替换占位资源。')
    add_table(tmp, ['阶段', '开发内容', '通过条件'], STEPS, [3.0, 7.1, 6.3])
    tmp.add_paragraph('P0 通过后，P1 再扩展 AI 密封竞拍、动态市场与完整养成；P2 再规划真人交易、商业化和运营。')
    blocks = [b for b in tmp.element.body if b.tag != qn('w:sectPr')]
    for index, block in enumerate(blocks):
        body.insert(index, block)

def normalize(doc):
    for sec in doc.sections:
        sec.top_margin = Cm(1.8)
        sec.bottom_margin = Cm(1.8)
        sec.left_margin = Cm(2.0)
        sec.right_margin = Cm(2.0)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
        if name == 'Normal':
            style.font.size = Pt(10.5)
            style.paragraph_format.space_after = Pt(5)
            style.paragraph_format.line_spacing = 1.2
        elif name == 'Title':
            style.font.size = Pt(21)
        elif name.startswith('Heading'):
            style.font.size = Pt(14 if name == 'Heading 1' else 11.5)
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(10)
            style.paragraph_format.space_after = Pt(5)
    for parent in [doc.styles.element, doc.element]:
        for border in list(parent.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    for p in paragraphs(doc):
        # Remove inherited manual page breaks introduced for old paragraph lengths.
        p.paragraph_format.page_break_before = False
        heading = p.style.name.startswith('Heading') or p.style.name == 'Title'
        if heading:
            p.paragraph_format.keep_with_next = True
        for r in p.runs:
            r.font.name = 'Microsoft YaHei'
            r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
            if heading:
                r.font.color.rgb = RGBColor(0, 0, 0)
                r.font.underline = False
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        prop = table._tbl.tblPr
        for old in list(prop.findall(qn('w:tblBorders'))):
            prop.remove(old)
        borders = OxmlElement('w:tblBorders')
        for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
            b = OxmlElement('w:' + side)
            for k, v in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                b.set(qn('w:' + k), v)
            borders.append(b)
        prop.append(borders)
        for index, row in enumerate(table.rows):
            row.height = None
            trpr = row._tr.get_or_add_trPr()
            for x in list(trpr.findall(qn('w:trHeight'))) + list(trpr.findall(qn('w:tblHeader'))) + list(trpr.findall(qn('w:cantSplit'))):
                trpr.remove(x)
            no_split = OxmlElement('w:cantSplit')
            no_split.set(qn('w:val'), 'true')
            trpr.append(no_split)
            if index == 0:
                repeat = OxmlElement('w:tblHeader')
                repeat.set(qn('w:val'), 'true')
                trpr.append(repeat)
            for c in row.cells:
                c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp = c._tc.get_or_add_tcPr()
                for old in list(cp.findall(qn('w:shd'))) + list(cp.findall(qn('w:tcMar'))):
                    cp.remove(old)
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'DCE6F1' if index == 0 else 'FFFFFF'); cp.append(shade)
                margin = OxmlElement('w:tcMar')
                for side, value in [('top', '90'), ('bottom', '90'), ('left', '110'), ('right', '110')]:
                    m = OxmlElement('w:' + side); m.set(qn('w:w'), value); m.set(qn('w:type'), 'dxa'); margin.append(m)
                cp.append(margin)
                for p in c.paragraphs:
                    p.paragraph_format.keep_with_next = index == 0
                    p.paragraph_format.keep_together = True
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.line_spacing = 1.15
                    for r in p.runs:
                        r.font.size = Pt(9.5)
                        r.font.color.rgb = RGBColor(0, 0, 0)
                        r.font.bold = index == 0

def finish(doc, role):
    normalize(doc)
    path = ROOT / FILES[role][1]
    doc.save(path)
    return path

if __name__ == '__main__':
    # Read-only source/game fingerprints for final verification.
    manifest = {'sources': {}, 'code': {}}
    for role, names in FILES.items():
        path = SOURCE / names[0]
        manifest['sources'][str(path)] = hashlib.sha256(path.read_bytes()).hexdigest()
    repo = Path('F:/HunyueRoblox/TravelBuddy')
    for name in ['src', 'scripts', 'tests', 'design']:
        for path in (repo / name).rglob('*'):
            if path.is_file() and '__pycache__' not in path.parts:
                manifest['code'][str(path)] = hashlib.sha256(path.read_bytes()).hexdigest()
    (QA / 'before.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print('Recorded sources and game project fingerprints', len(manifest['code']))
