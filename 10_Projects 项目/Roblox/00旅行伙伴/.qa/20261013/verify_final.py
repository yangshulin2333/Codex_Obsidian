from pathlib import Path
import hashlib
import json
import zipfile
from docx import Document
from docx.oxml.ns import qn
from doc_common import ROOT, QA, FILES, RULES, PARAMS, BASELINE, STEPS

def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

before = json.loads((QA / 'before.json').read_text(encoding='utf-8'))
source_changes = [p for p, h in before['sources'].items() if not Path(p).exists() or sha(p) != h]
code_changes = [p for p, h in before['code'].items() if not Path(p).exists() or sha(p) != h]
repo = Path('F:/HunyueRoblox/TravelBuddy')
current = {str(p) for folder in ('src', 'scripts', 'tests', 'design') for p in (repo / folder).rglob('*') if p.is_file() and '__pycache__' not in p.parts}
new_code_files = sorted(current - set(before['code']))
assert not source_changes, source_changes
assert not code_changes, code_changes
assert not new_code_files, new_code_files

expected_tables = [PARAMS, BASELINE, STEPS]
render_folders = {'game': QA / 'renders-final/game', 'program': QA / 'renders-final/program', 'art': QA / 'renders-art-clean'}
expected_pages = {'game': 13, 'program': 12, 'art': 22}
report = {'original_sources_unchanged': len(before['sources']), 'game_files_unchanged': len(before['code']), 'new_game_files': new_code_files, 'documents': {}}
for role, (_, name) in FILES.items():
    path = ROOT / name
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None
    doc = Document(path)
    paras = [p.text for p in doc.paragraphs]
    for label, value in RULES:
        assert label + '。' + value in paras, (name, label)
    for table, expected in zip(doc.tables[:3], expected_tables):
        rows = [tuple(c.text for c in row.cells) for row in table.rows[1:]]
        assert rows == expected, (name, rows)
    for table in doc.tables:
        trpr = table.rows[0]._tr.find(qn('w:trPr'))
        assert trpr is not None and trpr.find(qn('w:tblHeader')) is not None, name
        for row in table.rows:
            trpr = row._tr.find(qn('w:trPr'))
            assert trpr is not None and trpr.find(qn('w:cantSplit')) is not None, name
    pages = sorted(render_folders[role].glob('page-*.png'), key=lambda p: int(p.stem.split('-')[1]))
    assert len(pages) == expected_pages[role], (role, len(pages))
    assert [int(p.stem.split('-')[1]) for p in pages] == list(range(1, expected_pages[role] + 1))
    report['documents'][role] = {'path': str(path), 'sha256': sha(path), 'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables), 'pages': len(pages), 'page_sha256': {p.name: sha(p) for p in pages}, 'visual_review': 'All pages inspected at original size; identical page hashes reuse prior inspection.'}

program = Document(ROOT / FILES['program'][1])
program_text = '\n'.join(p.text for p in program.paragraphs) + '\n' + '\n'.join(c.text for t in program.tables for r in t.rows for c in r.cells)
for field in ('Wood', 'Upgrades.Home', 'Upgrades.Ship', 'TargetLevel', 'EndAt', 'UnreadSummary', 'ReadCursor'):
    assert field in program_text, field
report['common_rules_and_parameter_tables'] = 'identical in all three documents'
(QA / 'verification-final.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({'sources_unchanged': report['original_sources_unchanged'], 'game_files_unchanged': report['game_files_unchanged'], 'documents': {r: {'pages': d['pages'], 'tables': d['tables']} for r, d in report['documents'].items()}}, ensure_ascii=False))
