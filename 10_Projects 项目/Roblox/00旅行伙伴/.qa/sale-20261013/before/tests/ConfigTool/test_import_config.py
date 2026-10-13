from __future__ import annotations

import contextlib
import copy
import importlib.util
import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest import mock

from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink


ROOT = Path(__file__).resolve().parents[2]
IMPORTER_PATH = ROOT / "scripts/config/import_config.py"
SPEC = importlib.util.spec_from_file_location("import_config", IMPORTER_PATH)
assert SPEC is not None and SPEC.loader is not None
import_config = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = import_config
SPEC.loader.exec_module(import_config)


def make_model(schema):
    sheets = {}
    for sheet in schema["sheets"]:
        sheets[sheet["name"]] = {
            "descriptions": [column["descriptionZh"] for column in sheet["columns"]],
            "keys": [column["key"] for column in sheet["columns"]],
            "types": [column["type"] for column in sheet["columns"]],
            "rows": [],
            "formulas": set(),
            "workbook": sheet["workbook"],
        }

    sheets["ExampleGroups"]["rows"] = [
        {"row": 4, "values": ["Basic", "基础组"]},
        {"row": 5, "values": ["Advanced", "进阶组"]},
    ]
    item_values = [
        dict(Id="BasicItem", DisplayName="基础示例", GroupId="Basic", Tier="Common",
             Price=100, Enabled=True, Tags="starter,shop", Note=None),
        dict(Id="AdvancedItem", DisplayName="进阶示例", GroupId="Advanced", Tier="Rare",
             Price=500.5, Enabled=False, Tags="rare,shop", Note="仅用于验证"),
    ]
    sheets["ExampleItems"]["rows"] = [
        {"row": row, "values": [values[key] for key in sheets["ExampleItems"]["keys"]]}
        for row, values in enumerate(item_values, 4)
    ]
    sheets["ServerSettings"]["rows"] = [
        {"row": 4, "values": ["RewardMultiplier", 1, True]},
        {"row": 5, "values": ["SpawnInterval", 5, True]},
    ]
    return {"sheetNames": list(sheets), "sheets": sheets}


class ImportConfigTest(unittest.TestCase):
    def test_asset_id_normalizes_without_numeric_precision_loss(self):
        column = dict(key="AssetId", type="string", required=True, format="assetId")
        for raw in ("12345678901234567890", "rbxassetid://12345678901234567890"):
            value, issues = import_config.validate_value(raw, column, "Animations", 4, 2)
            self.assertEqual([], issues)
            self.assertEqual("rbxassetid://12345678901234567890", value)

    def test_asset_id_rejects_bad_text_numeric_cells_and_blank(self):
        column = dict(key="AssetId", type="string", required=True, format="assetId")
        for raw in ("0", "-1", "1.5", "1e15", "https://roblox.com/123", "rbxassetid://", "１２３", 123, True, None):
            with self.subTest(raw=raw):
                _, issues = import_config.validate_value(raw, column, "Animations", 4, 2)
                self.assertTrue(issues)
                self.assertIn("Cell=B4", issues[0].format())

    def test_asset_tables_reject_duplicate_keys_and_export_content_ids(self):
        model = make_model(self.schema)
        model["sheets"]["Images"]["rows"] = [
            {"row": 4, "values": ["SampleIcon", "123"]},
            {"row": 5, "values": ["SampleIcon", "456"]},
        ]
        _, issues = self.validate(model)
        self.assertTrue(issues)
        model["sheets"]["Images"]["rows"].pop()
        rows, issues = self.validate(model)
        self.assertEqual([], issues)
        self.assertEqual("rbxassetid://123", rows["Images"][0]["AssetId"])

    def test_asset_format_rejects_wrong_schema_type(self):
        schema = copy.deepcopy(self.schema)
        schema["sheets"][0]["columns"][1].update(type="number", format="assetId")
        with self.assertRaises(import_config.ConfigFailure):
            import_config.validate_schema(schema)

    def test_optional_image_asset_ids_allow_pending_records_without_fake_ids(self):
        model = make_model(self.schema)
        model["sheets"]["Images"]["rows"] = [
            {"row": 4, "values": ["PendingIcon", None]},
        ]
        rows, issues = self.validate(model)
        self.assertEqual([], issues)
        output = import_config.build_outputs(self.schema, rows)["Shared"]["Images.lua"]
        self.assertIn('["Id"] = "PendingIcon"', output)
        self.assertNotIn('["AssetId"] =', output)

        column = next(sheet for sheet in self.schema["sheets"] if sheet["name"] == "Images")["columns"][1]
        for raw in (None, "", "  "):
            with self.subTest(raw=raw):
                value, issues = import_config.validate_value(raw, column, "Images", 4, 2)
                self.assertIsNone(value)
                self.assertEqual([], issues)
        for raw in ("pending", "0", "rbxassetid://", "https://example.com/123", 123, True):
            with self.subTest(raw=raw):
                _, issues = import_config.validate_value(raw, column, "Images", 4, 2)
                self.assertTrue(issues)

    def test_export_false_validates_fields_without_emitting_them(self):
        schema = copy.deepcopy(self.schema)
        item_sheet = next(sheet for sheet in schema["sheets"] if sheet["name"] == "ExampleItems")
        price = next(column for column in item_sheet["columns"] if column["key"] == "Price")
        price["export"] = False
        model = make_model(schema)
        rows, issues = import_config.validate_workbook_model(model, schema)
        self.assertEqual([], issues)
        output = import_config.build_outputs(schema, rows)["Shared"]["ExampleItems.lua"]
        self.assertNotIn('["Price"] =', output)
        self.assertNotIn('["Note"] =', output)
        self.assertNotIn("仅用于验证", output)
        self.assertIn('["Tags"] = { "rare", "shop" }', output)

        price_index = model["sheets"]["ExampleItems"]["keys"].index("Price")
        model["sheets"]["ExampleItems"]["rows"][0]["values"][price_index] = "wrong"
        _, issues = import_config.validate_workbook_model(model, schema)
        self.assertIn("CONFIG_TYPE_INVALID", {issue.code for issue in issues})

    def test_export_and_integer_schema_flags_are_typed_constraints(self):
        for update in (
            {"export": "false"}, {"export": 0}, {"integer": "true"},
            {"integer": 1}, {"integer": True, "type": "string"},
        ):
            with self.subTest(update=update):
                schema = copy.deepcopy(self.schema)
                schema["sheets"][2]["columns"][1].update(update)
                with self.assertRaises(import_config.ConfigFailure) as context:
                    import_config.validate_schema(schema)
                self.assertIn("CONFIG_SCHEMA_INVALID", {issue.code for issue in context.exception.issues})

    def test_integer_fields_reject_fractions_without_restricting_decimal_fields(self):
        column = dict(key="Count", type="number", required=True, integer=True)
        for raw in (0, -2, 3.0):
            with self.subTest(raw=raw):
                value, issues = import_config.validate_value(raw, column, "Records", 4, 2)
                self.assertEqual([], issues)
                self.assertEqual(raw, value)
        _, issues = import_config.validate_value(2.5, column, "Records", 4, 2)
        self.assertIn("CONFIG_INTEGER_REQUIRED", {issue.code for issue in issues})
        _, issues = import_config.validate_value(True, column, "Records", 4, 2)
        self.assertIn("CONFIG_TYPE_INVALID", {issue.code for issue in issues})
        column.pop("integer")
        value, issues = import_config.validate_value(2.5, column, "Records", 4, 2)
        self.assertEqual([], issues)
        self.assertEqual(2.5, value)

    @classmethod
    def setUpClass(cls):
        cls.schema = json.loads(
            (ROOT / "tests/ConfigTool/fixtures/schema.json").read_text(encoding="utf-8")
        )
        import_config.validate_schema(cls.schema)

    def validate(self, model):
        return import_config.validate_workbook_model(model, self.schema)

    def test_valid_model_generates_shared_and_server_outputs(self):
        typed_rows, issues = self.validate(make_model(self.schema))
        self.assertEqual([], issues)

        outputs = import_config.build_outputs(self.schema, typed_rows)
        for scope in ("Shared", "Server"):
            expected = {
                f"{sheet['name']}.lua"
                for sheet in self.schema["sheets"]
                if sheet["scope"] == scope
            }
            self.assertEqual(expected, set(outputs[scope]))
        self.assertIn('["GroupId"] = "Basic"', outputs["Shared"]["ExampleItems.lua"])
        self.assertIn(
            '["Tags"] = { "starter", "shop" }', outputs["Shared"]["ExampleItems.lua"]
        )
        for scope_outputs in outputs.values():
            for content in scope_outputs.values():
                self.assertIn(
                    "-- 此文件由 Excel Config 工具自动生成，请勿手动修改。",
                    content,
                )

    def test_formatted_error_includes_sheet_and_cell(self):
        model = make_model(self.schema)
        model["sheets"]["ExampleItems"]["rows"][0]["values"][1] = None
        _, issues = self.validate(model)
        issue = next(
            issue for issue in issues if issue.code == "CONFIG_REQUIRED_MISSING"
        )
        self.assertEqual(
            "CONFIG_REQUIRED_MISSING | Sheet=ExampleItems | Cell=B4 | "
            "字段 DisplayName 为必填项",
            issue.format(),
        )

    def test_description_must_match_schema(self):
        model = make_model(self.schema)
        model["sheets"]["ExampleItems"]["descriptions"][0] = "错误说明"
        _, issues = self.validate(model)
        self.assertIn("CONFIG_DESCRIPTION_MISMATCH", {issue.code for issue in issues})

    def test_type_header_must_match_schema(self):
        model = make_model(self.schema)
        model["sheets"]["ExampleItems"]["types"][4] = "string"
        _, issues = self.validate(model)
        self.assertIn("CONFIG_TYPE_HEADER_MISMATCH", {issue.code for issue in issues})

    def test_formula_is_rejected(self):
        model = make_model(self.schema)
        model["sheets"]["ExampleItems"]["formulas"].add((3, 5))
        _, issues = self.validate(model)
        self.assertIn("CONFIG_FORMULA_NOT_ALLOWED", {issue.code for issue in issues})

    def test_duplicate_id_is_rejected(self):
        model = make_model(self.schema)
        model["sheets"]["ExampleItems"]["rows"][1]["values"][0] = "BasicItem"
        _, issues = self.validate(model)
        self.assertIn("CONFIG_DUPLICATE_VALUE", {issue.code for issue in issues})

    def test_missing_required_value_is_rejected(self):
        model = make_model(self.schema)
        model["sheets"]["ExampleItems"]["rows"][0]["values"][1] = None
        _, issues = self.validate(model)
        self.assertIn("CONFIG_REQUIRED_MISSING", {issue.code for issue in issues})

    def test_wrong_type_is_rejected(self):
        model = make_model(self.schema)
        model["sheets"]["ExampleItems"]["rows"][0]["values"][5] = "yes"
        _, issues = self.validate(model)
        self.assertIn("CONFIG_TYPE_INVALID", {issue.code for issue in issues})

    def test_excel_true_text_is_accepted(self):
        model = make_model(self.schema)
        model["sheets"]["ExampleItems"]["rows"][0]["values"][5] = "TRUE"
        typed_rows, issues = self.validate(model)
        self.assertEqual([], issues)
        self.assertIs(typed_rows["ExampleItems"][0]["Enabled"], True)

    def test_number_array_is_parsed(self):
        value, issues = import_config.validate_value(
            "1, 2.5, 3",
            {"key": "Weights", "type": "number[]", "required": True},
            "ExampleItems",
            4,
            1,
        )
        self.assertEqual([], issues)
        self.assertEqual([1, 2.5, 3], value)

    def test_invalid_number_array_is_rejected(self):
        _, issues = import_config.validate_value(
            "1, wrong, 3",
            {"key": "Weights", "type": "number[]", "required": True},
            "ExampleItems",
            4,
            1,
        )
        self.assertIn("CONFIG_TYPE_INVALID", {issue.code for issue in issues})

    def test_number_below_minimum_is_rejected(self):
        model = make_model(self.schema)
        model["sheets"]["ExampleItems"]["rows"][0]["values"][4] = -1
        _, issues = self.validate(model)
        self.assertIn("CONFIG_NUMBER_OUT_OF_RANGE", {issue.code for issue in issues})

    def test_unknown_enum_is_rejected(self):
        model = make_model(self.schema)
        model["sheets"]["ExampleItems"]["rows"][0]["values"][3] = "Unknown"
        _, issues = self.validate(model)
        self.assertIn("CONFIG_ENUM_INVALID", {issue.code for issue in issues})

    def test_missing_reference_is_rejected(self):
        model = make_model(self.schema)
        model["sheets"]["ExampleItems"]["rows"][0]["values"][2] = "Missing"
        _, issues = self.validate(model)
        self.assertIn("CONFIG_REFERENCE_NOT_FOUND", {issue.code for issue in issues})

    def test_unknown_sheet_is_rejected(self):
        model = make_model(self.schema)
        model["sheetNames"].append("UnknownSheet")
        model["sheets"]["UnknownSheet"] = {
            "descriptions": [],
            "keys": [],
            "types": [],
            "rows": [],
            "formulas": set(),
            "workbook": "GameConfig.xlsx",
        }
        _, issues = self.validate(model)
        self.assertIn("CONFIG_UNKNOWN_SHEET", {issue.code for issue in issues})

    def test_missing_sheet_is_rejected(self):
        model = make_model(self.schema)
        model["sheetNames"].remove("ServerSettings")
        del model["sheets"]["ServerSettings"]
        _, issues = self.validate(model)
        self.assertIn("CONFIG_SHEET_MISSING", {issue.code for issue in issues})

    def test_unknown_column_is_rejected(self):
        model = make_model(self.schema)
        model["sheets"]["ExampleGroups"]["descriptions"].append("未知字段")
        model["sheets"]["ExampleGroups"]["keys"].append("Unknown")
        model["sheets"]["ExampleGroups"]["types"].append("string")
        _, issues = self.validate(model)
        self.assertIn("CONFIG_UNKNOWN_COLUMN", {issue.code for issue in issues})

    def test_data_under_blank_unknown_column_is_rejected(self):
        model = make_model(self.schema)
        model["sheets"]["ExampleGroups"]["descriptions"].append(None)
        model["sheets"]["ExampleGroups"]["keys"].append(None)
        model["sheets"]["ExampleGroups"]["types"].append(None)
        model["sheets"]["ExampleGroups"]["rows"][0]["values"].append("unexpected")
        _, issues = self.validate(model)
        self.assertIn("CONFIG_UNKNOWN_COLUMN", {issue.code for issue in issues})

    def test_number_limit_on_string_is_invalid_schema(self):
        schema = copy.deepcopy(self.schema)
        schema["sheets"][0]["columns"][0]["min"] = 0
        with self.assertRaises(import_config.ConfigFailure) as context:
            import_config.validate_schema(schema)
        self.assertIn(
            "CONFIG_SCHEMA_INVALID", {issue.code for issue in context.exception.issues}
        )

    def test_real_workbook_matches_schema(self):
        schema = import_config.load_schema(ROOT / "design/config/config-schema.json")
        model = import_config.read_workbooks_model(
            ROOT / "design/config/workbooks", schema["layout"]
        )
        _, issues = import_config.validate_workbook_model(model, schema)
        self.assertEqual([], issues)

    def test_sync_creates_missing_workbook_sheet_and_headers(self):
        schema = copy.deepcopy(self.schema)
        schema["sheets"] = [
            sheet for sheet in schema["sheets"] if sheet["name"] == "ServerSettings"
        ]

        with tempfile.TemporaryDirectory() as directory:
            workbook_dir = Path(directory)
            changes = import_config.sync_workbooks_from_schema(workbook_dir, schema)

            self.assertEqual(1, len(changes))
            path = workbook_dir / "GameConfig.xlsx"
            workbook = import_config.load_workbook(path)
            try:
                self.assertEqual(["ServerSettings"], workbook.sheetnames)
                worksheet = workbook["ServerSettings"]
                self.assertEqual(schema["sheets"][0]["columns"][0]["descriptionZh"], worksheet["A1"].value)
                self.assertEqual(["Id", "NumberValue", "Enabled"], [
                    worksheet.cell(2, column).value for column in range(1, 4)
                ])
                self.assertEqual(["string", "number", "boolean"], [
                    worksheet.cell(3, column).value for column in range(1, 4)
                ])
                self.assertEqual("A4", str(worksheet.freeze_panes))
            finally:
                workbook.close()

    def test_sync_is_idempotent_and_preserves_existing_data(self):
        schema = copy.deepcopy(self.schema)
        schema["sheets"] = [
            sheet for sheet in schema["sheets"] if sheet["name"] == "ServerSettings"
        ]

        with tempfile.TemporaryDirectory() as directory:
            workbook_dir = Path(directory)
            import_config.sync_workbooks_from_schema(workbook_dir, schema)
            path = workbook_dir / "GameConfig.xlsx"
            workbook = import_config.load_workbook(path)
            workbook["ServerSettings"]["A4"] = "KeepMe"
            workbook.save(path)
            workbook.close()

            changes = import_config.sync_workbooks_from_schema(workbook_dir, schema)
            workbook = import_config.load_workbook(path)
            try:
                self.assertEqual([], changes)
                self.assertEqual("KeepMe", workbook["ServerSettings"]["A4"].value)
            finally:
                workbook.close()

    def test_sync_only_appends_new_trailing_columns(self):
        schema = copy.deepcopy(self.schema)
        schema["sheets"] = [
            sheet for sheet in schema["sheets"] if sheet["name"] == "ServerSettings"
        ]

        with tempfile.TemporaryDirectory() as directory:
            workbook_dir = Path(directory)
            path = workbook_dir / "GameConfig.xlsx"
            workbook = import_config.Workbook()
            worksheet = workbook.active
            worksheet.title = "ServerSettings"
            first_column = schema["sheets"][0]["columns"][0]
            worksheet["A1"] = first_column["descriptionZh"]
            worksheet["A2"] = first_column["key"]
            worksheet["A3"] = first_column["type"]
            worksheet["A4"] = "KeepMe"
            workbook.save(path)
            workbook.close()

            changes = import_config.sync_workbooks_from_schema(workbook_dir, schema)
            workbook = import_config.load_workbook(path)
            try:
                self.assertEqual(1, len(changes))
                self.assertIn("Action=APPEND_COLUMNS", changes[0])
                self.assertEqual("KeepMe", workbook["ServerSettings"]["A4"].value)
                self.assertEqual("NumberValue", workbook["ServerSettings"]["B2"].value)
                self.assertEqual("Enabled", workbook["ServerSettings"]["C2"].value)
            finally:
                workbook.close()

    def test_sync_does_not_rewrite_conflicting_existing_headers(self):
        schema = copy.deepcopy(self.schema)
        schema["sheets"] = [
            sheet for sheet in schema["sheets"] if sheet["name"] == "ServerSettings"
        ]

        with tempfile.TemporaryDirectory() as directory:
            workbook_dir = Path(directory)
            path = workbook_dir / "GameConfig.xlsx"
            workbook = import_config.Workbook()
            worksheet = workbook.active
            worksheet.title = "ServerSettings"
            worksheet["A2"] = "WrongKey"
            worksheet["A4"] = "KeepMe"
            workbook.save(path)
            workbook.close()

            changes = import_config.sync_workbooks_from_schema(workbook_dir, schema)
            workbook = import_config.load_workbook(path)
            try:
                self.assertEqual([], changes)
                self.assertEqual("WrongKey", workbook["ServerSettings"]["A2"].value)
                self.assertEqual("KeepMe", workbook["ServerSettings"]["A4"].value)
            finally:
                workbook.close()

    def test_sync_preserves_drafts_without_english_headers(self):
        schema = copy.deepcopy(self.schema)
        schema["sheets"] = schema["sheets"][:1]
        sheet = schema["sheets"][0]
        for cell in ("A1", "A3", "A4"):
            with self.subTest(cell=cell), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / sheet["workbook"]
                workbook = import_config.Workbook()
                worksheet = workbook.active
                worksheet.title = sheet["name"]
                worksheet[cell] = "KeepDraft"
                workbook.save(path)
                workbook.close()
                before = path.read_bytes()
                self.assertEqual([], import_config.sync_workbooks_from_schema(Path(directory), schema))
                self.assertEqual(before, path.read_bytes())
                model = import_config.read_workbooks_model(Path(directory), schema["layout"])
                _, issues = import_config.validate_workbook_model(model, schema)
                self.assertTrue(issues)

    def test_sync_preserves_occupied_trailing_columns(self):
        schema = copy.deepcopy(self.schema)
        schema["sheets"] = schema["sheets"][:1]
        sheet = schema["sheets"][0]
        for cell in ("B1", "B3", "B4"):
            with self.subTest(cell=cell), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / sheet["workbook"]
                workbook = import_config.Workbook()
                worksheet = workbook.active
                worksheet.title = sheet["name"]
                column = sheet["columns"][0]
                worksheet["A1"] = column["descriptionZh"]
                worksheet["A2"] = column["key"]
                worksheet["A3"] = column["type"]
                worksheet[cell] = "KeepDraft"
                workbook.save(path)
                workbook.close()
                before = path.read_bytes()
                self.assertEqual([], import_config.sync_workbooks_from_schema(Path(directory), schema))
                self.assertEqual(before, path.read_bytes())
                model = import_config.read_workbooks_model(Path(directory), schema["layout"])
                _, issues = import_config.validate_workbook_model(model, schema)
                self.assertTrue(issues)

    def test_sync_write_failure_keeps_original_workbook(self):
        schema = copy.deepcopy(self.schema)
        schema["sheets"] = [
            sheet for sheet in schema["sheets"] if sheet["name"] == "ServerSettings"
        ]

        with tempfile.TemporaryDirectory() as directory:
            workbook_dir = Path(directory)
            path = workbook_dir / "GameConfig.xlsx"
            workbook = import_config.Workbook()
            workbook.active.title = "Existing"
            workbook.active["A1"] = "KeepMe"
            workbook.save(path)
            workbook.close()
            original = path.read_bytes()

            with mock.patch.object(
                import_config.os,
                "replace",
                side_effect=PermissionError("workbook locked"),
            ):
                with self.assertRaises(import_config.ConfigFailure) as context:
                    import_config.sync_workbooks_from_schema(workbook_dir, schema)

            self.assertIn(
                "CONFIG_WORKBOOK_WRITE_FAILED",
                {issue.code for issue in context.exception.issues},
            )
            self.assertEqual(original, path.read_bytes())
            self.assertEqual([], list(workbook_dir.glob(".*.config-sync-*.xlsx")))

    def test_run_import_syncs_then_generates(self):
        schema = copy.deepcopy(self.schema)
        schema["sheets"] = [
            sheet for sheet in schema["sheets"] if sheet["name"] == "ServerSettings"
        ]

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_dir = root / "workbooks"
            schema_path = root / "config-schema.json"
            schema_path.write_text(json.dumps(schema), encoding="utf-8")
            args = SimpleNamespace(
                repo_root=str(root),
                workbook_dir=str(workbook_dir),
                schema=str(schema_path),
                check=False,
            )

            output = io.StringIO()
            with contextlib.redirect_stdout(output):
                self.assertEqual(0, import_config.run(args))

            self.assertIn("Action=CREATE_SHEET", output.getvalue())
            self.assertIn("CONFIG_IMPORT_OK", output.getvalue())
            self.assertTrue((workbook_dir / "GameConfig.xlsx").is_file())
            self.assertTrue(
                (
                    import_config.output_targets(root)["Server"] / "ServerSettings.lua"
                ).is_file()
            )

    def test_check_mode_never_creates_workbook_structure(self):
        schema = copy.deepcopy(self.schema)
        schema["sheets"] = [
            sheet for sheet in schema["sheets"] if sheet["name"] == "ServerSettings"
        ]

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workbook_dir = root / "workbooks"
            workbook_dir.mkdir()
            schema_path = root / "config-schema.json"
            schema_path.write_text(json.dumps(schema), encoding="utf-8")
            args = SimpleNamespace(
                repo_root=str(root),
                workbook_dir=str(workbook_dir),
                schema=str(schema_path),
                check=True,
            )

            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(1, import_config.run(args))
            self.assertFalse((workbook_dir / "GameConfig.xlsx").exists())

    def test_multiple_workbooks_are_merged(self):
        with tempfile.TemporaryDirectory() as directory:
            workbook_dir = Path(directory)
            first_path = workbook_dir / "First.xlsx"
            second_path = workbook_dir / "Second.xlsx"
            first_path.touch()
            second_path.touch()

            def fake_read(path, _layout):
                sheet_name = path.stem
                return {
                    "sheetNames": [sheet_name],
                    "sheets": {
                        sheet_name: {
                            "descriptions": [],
                            "keys": [],
                            "types": [],
                            "rows": [],
                            "formulas": set(),
                            "workbook": path.name,
                        }
                    },
                }

            with mock.patch.object(import_config, "read_workbook_model", side_effect=fake_read):
                model = import_config.read_workbooks_model(
                    workbook_dir, self.schema["layout"]
                )

            self.assertEqual(["First", "Second"], model["sheetNames"])

    def test_sheet_in_wrong_workbook_is_rejected(self):
        model = make_model(self.schema)
        model["sheets"]["ExampleItems"]["workbook"] = "Other.xlsx"
        _, issues = self.validate(model)
        self.assertIn("CONFIG_SHEET_WRONG_WORKBOOK", {issue.code for issue in issues})

    def test_generation_is_deterministic(self):
        typed_rows, issues = self.validate(make_model(self.schema))
        self.assertEqual([], issues)
        first = import_config.build_outputs(self.schema, typed_rows)
        second = import_config.build_outputs(self.schema, copy.deepcopy(typed_rows))
        self.assertEqual(first, second)

    def test_atomic_write_matches_check_mode(self):
        typed_rows, issues = self.validate(make_model(self.schema))
        self.assertEqual([], issues)
        outputs = import_config.build_outputs(self.schema, typed_rows)

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            import_config.write_outputs_atomic(root, outputs)
            self.assertEqual([], import_config.check_outputs(root, outputs))

            shared = import_config.output_targets(root)["Shared"] / "ExampleItems.lua"
            shared.write_text("outdated", encoding="utf-8")
            self.assertIn(
                "CONFIG_OUTPUT_OUTDATED",
                {issue.code for issue in import_config.check_outputs(root, outputs)},
            )

    def test_atomic_write_keeps_backup_outside_source_tree(self):
        typed_rows, issues = self.validate(make_model(self.schema))
        self.assertEqual([], issues)
        outputs = import_config.build_outputs(self.schema, typed_rows)

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            import_config.write_outputs_atomic(root, outputs)
            destinations = []
            real_replace = import_config.os.replace

            def record_replace(source, destination):
                destinations.append(Path(destination))
                return real_replace(source, destination)

            with mock.patch.object(
                import_config.os, "replace", side_effect=record_replace
            ):
                import_config.write_outputs_atomic(root, outputs)

            backup_destinations = [
                path for path in destinations if path.name.endswith(".previous")
            ]
            self.assertEqual(2, len(backup_destinations))
            self.assertTrue(
                all(path.parent.name.startswith(".config-build-") for path in backup_destinations)
            )
            self.assertFalse(
                any("config-backup" in path.name for path in destinations)
            )

    def test_atomic_write_cleanup_failure_keeps_committed_outputs(self):
        old_outputs = {
            "Shared": {"OldShared.lua": "return { Old = true }\n"},
            "Server": {"OldServer.lua": "return { Old = true }\n"},
        }
        new_outputs = {
            "Shared": {"NewShared.lua": "return { New = true }\n"},
            "Server": {"NewServer.lua": "return { New = true }\n"},
        }

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            import_config.write_outputs_atomic(root, old_outputs)
            real_rmtree = import_config.shutil.rmtree
            failure_injected = False

            def fail_second_backup_once(path, *args, **kwargs):
                nonlocal failure_injected
                if Path(path).name == "Server.previous" and not failure_injected:
                    failure_injected = True
                    raise OSError("simulated backup cleanup failure")
                return real_rmtree(path, *args, **kwargs)

            with mock.patch.object(
                import_config.shutil,
                "rmtree",
                side_effect=fail_second_backup_once,
            ):
                import_config.write_outputs_atomic(root, new_outputs)

            self.assertTrue(failure_injected)
            for scope, target in import_config.output_targets(root).items():
                self.assertEqual(
                    set(new_outputs[scope]),
                    {path.name for path in target.iterdir()},
                )
                for file_name, content in new_outputs[scope].items():
                    self.assertEqual(
                        content,
                        (target / file_name).read_text(encoding="utf-8"),
                    )

    def test_validation_failure_keeps_previous_outputs(self):
        model = make_model(self.schema)
        model["sheets"]["ExampleItems"]["rows"][1]["values"][0] = "BasicItem"

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            schema_path = root / "config-schema.json"
            schema_path.write_text(json.dumps(self.schema), encoding="utf-8")
            targets = import_config.output_targets(root)
            for target in targets.values():
                target.mkdir(parents=True)
                (target / "Previous.lua").write_text("previous", encoding="utf-8")

            args = SimpleNamespace(
                repo_root=str(root),
                workbook_dir=str(root / "workbooks"),
                schema=str(schema_path),
                check=False,
            )
            with mock.patch.object(import_config, "read_workbooks_model", return_value=model):
                with contextlib.redirect_stdout(io.StringIO()):
                    self.assertEqual(1, import_config.run(args))

            for target in targets.values():
                self.assertEqual("previous", (target / "Previous.lua").read_text(encoding="utf-8"))


def note_schema(keys=("Id", "Value", "Extra", "Note")):
    columns = []
    for key in keys:
        column = dict(key=key, descriptionZh=f"示例字段 {key}", type="string", required=False)
        if key == "Id":
            column.update(required=True, unique=True)
        if key == "Count":
            column.update(type="number", integer=True)
        if key == "Note":
            column["export"] = False
        columns.append(column)
    return {
        "version": 1,
        "layout": dict(descriptionRow=1, keyRow=2, typeRow=3, dataStartRow=4),
        "sheets": [dict(name="Records", workbook="ConfigTest.xlsx", descriptionZh="通用移列回归",
                        scope="Shared", primaryKey="Id", columns=columns)],
    }


def note_workbook(schema, keys):
    workbook = import_config.Workbook()
    worksheet = workbook.active
    worksheet.title = "Records"
    columns = {column["key"]: column for column in schema["sheets"][0]["columns"]}
    values = {"Id": "Keep", "Value": "original value", "Count": 4, "Note": "original note"}
    for index, key in enumerate(keys, 1):
        column = columns.get(key, dict(key=key, descriptionZh=f"示例字段 {key}", type="string"))
        worksheet.cell(1, index, column["descriptionZh"])
        worksheet.cell(2, index, key)
        worksheet.cell(3, index, column["type"])
        worksheet.cell(4, index, values.get(key, "keep unknown"))
    return workbook, worksheet


class NoteColumnTest(unittest.TestCase):
    def test_note_schema_requires_last_non_exported_column(self):
        schema = note_schema()
        import_config.validate_schema(schema)
        for export in (True, None):
            with self.subTest(export=export):
                invalid = copy.deepcopy(schema)
                note = invalid["sheets"][0]["columns"][-1]
                if export is None:
                    note.pop("export")
                else:
                    note["export"] = export
                with self.assertRaises(import_config.ConfigFailure) as context:
                    import_config.validate_schema(invalid)
                self.assertIn("CONFIG_SCHEMA_INVALID", {issue.code for issue in context.exception.issues})
        invalid = copy.deepcopy(schema)
        columns = invalid["sheets"][0]["columns"]
        columns.insert(1, columns.pop())
        with self.assertRaises(import_config.ConfigFailure):
            import_config.validate_schema(invalid)

    def test_validation_requires_note_last_even_before_automatic_sync(self):
        schema = note_schema(("Id", "Value", "Note"))
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "ConfigTest.xlsx"
            workbook, _ = note_workbook(schema, ["Id", "Note", "Value"])
            workbook.save(path)
            workbook.close()
            model = import_config.read_workbook_model(path, schema["layout"])
            _, issues = import_config.validate_workbook_model(model, schema)
            self.assertIn("CONFIG_HEADER_MISMATCH", {issue.code for issue in issues})

    def test_safe_note_move_and_insertion_preserve_cell_and_column_metadata(self):
        schema = note_schema()
        for keys in (["Id", "Note", "Value"], ["Id", "Value", "Note"]):
            with self.subTest(keys=keys), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "ConfigTest.xlsx"
                workbook, worksheet = note_workbook(schema, keys)
                note_index = keys.index("Note") + 1
                note = worksheet.cell(4, note_index)
                note.font = Font(bold=True, color="FF123456")
                note.fill = PatternFill("solid", fgColor="FFFFAA00")
                note.alignment = Alignment(wrap_text=True)
                note.number_format = "0.000"
                note.comment = Comment("keep comment", "Designer")
                note.hyperlink = "https://example.com/notes"
                dimension = worksheet.column_dimensions[note.column_letter]
                dimension.width, dimension.hidden, dimension.number_format = 29, True, "0.000"
                worksheet.cell(4, keys.index("Value") + 1).font = Font(italic=True)
                worksheet.freeze_panes = "A4"
                worksheet.auto_filter.ref = "A2:C4"
                worksheet.auto_filter.add_filter_column(note_index - 1, ["original note"])
                workbook.save(path)
                workbook.close()

                self.assertTrue(import_config.sync_workbooks_from_schema(Path(directory), schema))
                workbook = import_config.load_workbook(path)
                try:
                    worksheet = workbook["Records"]
                    self.assertEqual(["Id", "Value", "Extra", "Note"], [cell.value for cell in worksheet[2]])
                    self.assertEqual(["Keep", "original value", None, "original note"], [cell.value for cell in worksheet[4]])
                    self.assertTrue(worksheet["B4"].font.italic)
                    note = worksheet["D4"]
                    self.assertEqual("s", note.data_type)
                    self.assertTrue(note.font.bold)
                    self.assertEqual("FF123456", note.font.color.rgb)
                    self.assertEqual("FFFFAA00", note.fill.fgColor.rgb)
                    self.assertTrue(note.alignment.wrap_text)
                    self.assertEqual("0.000", note.number_format)
                    self.assertEqual("keep comment", note.comment.text)
                    self.assertEqual("Designer", note.comment.author)
                    self.assertEqual("https://example.com/notes", note.hyperlink.target)
                    self.assertEqual("D4", note.hyperlink.ref)
                    self.assertEqual(29, worksheet.column_dimensions["D"].width)
                    self.assertTrue(worksheet.column_dimensions["D"].hidden)
                    self.assertEqual("0.000", worksheet.column_dimensions["D"].number_format)
                    self.assertFalse(worksheet.column_dimensions["C"].hidden)
                    self.assertEqual(3, worksheet.auto_filter.filterColumn[0].colId)
                    self.assertEqual("A4", str(worksheet.freeze_panes))
                finally:
                    workbook.close()
                model = import_config.read_workbook_model(path, schema["layout"])
                rows, issues = import_config.validate_workbook_model(model, schema)
                self.assertEqual([], issues)
                self.assertEqual("original value", rows["Records"][0]["Value"])
                self.assertNotIn('["Note"] =', import_config.build_outputs(schema, rows)["Shared"]["Records.lua"])
                before = path.read_bytes()
                self.assertEqual([], import_config.sync_workbooks_from_schema(Path(directory), schema))
                self.assertEqual(before, path.read_bytes())

    def test_unnamed_nonempty_columns_block_note_sync_without_changing_file(self):
        schema = note_schema()
        for cell in ("D1", "D3", "D4"):
            with self.subTest(cell=cell), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "ConfigTest.xlsx"
                workbook, worksheet = note_workbook(schema, ["Id", "Note", "Value"])
                worksheet[cell] = "unnamed content"
                workbook.save(path)
                workbook.close()
                before = path.read_bytes()
                with self.assertRaises(import_config.ConfigFailure) as context:
                    import_config.sync_workbooks_from_schema(Path(directory), schema)
                self.assertIn("CONFIG_UNKNOWN_COLUMN", {issue.code for issue in context.exception.issues})
                self.assertEqual(before, path.read_bytes())

    def test_business_renames_and_reordering_are_not_inferred(self):
        schema = note_schema()
        for keys in (["Id", "Renamed", "Note"], ["Value", "Id", "Note"]):
            with self.subTest(keys=keys), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "ConfigTest.xlsx"
                workbook, _ = note_workbook(schema, keys)
                workbook.save(path)
                workbook.close()
                before = path.read_bytes()
                self.assertEqual([], import_config.sync_workbooks_from_schema(Path(directory), schema))
                self.assertEqual(before, path.read_bytes())
                model = import_config.read_workbook_model(path, schema["layout"])
                _, issues = import_config.validate_workbook_model(model, schema)
                self.assertIn("CONFIG_HEADER_MISMATCH", {issue.code for issue in issues})

    def test_complex_references_block_note_sync_and_keep_original_bytes(self):
        schema = note_schema()
        hazards = ("formula", "other_formula", "custom_validation", "forged_tool_rule",
                   "merged", "conditional_format", "sort", "other_chart", "internal_location", "internal_target")
        for hazard in hazards:
            with self.subTest(hazard=hazard), tempfile.TemporaryDirectory() as directory:
                path = Path(directory) / "ConfigTest.xlsx"
                workbook, worksheet = note_workbook(schema, ["Id", "Note", "Value"])
                if hazard == "formula":
                    worksheet["C4"] = "=1+1"
                elif hazard == "other_formula":
                    workbook.create_sheet("View")["A1"] = "=Records!B4"
                elif hazard in ("custom_validation", "forged_tool_rule"):
                    rule = DataValidation(type="custom", formula1="LEN(B4)>0")
                    if hazard == "forged_tool_rule":
                        rule.promptTitle, rule.errorTitle = "Note", "ConfigTool"
                    rule.add("B4:B1048576" if hazard == "forged_tool_rule" else "B4:B10")
                    worksheet.add_data_validation(rule)
                elif hazard == "merged":
                    worksheet.merge_cells("B4:C4")
                elif hazard == "conditional_format":
                    worksheet.conditional_formatting.add("B4", FormulaRule(formula=['B4<>""']))
                elif hazard == "sort":
                    worksheet.auto_filter.ref = "A2:C4"
                    worksheet.auto_filter.add_sort_condition("B4:B4")
                elif hazard == "other_chart":
                    chart = BarChart()
                    chart.add_data(Reference(worksheet, min_col=3, min_row=4, max_row=4))
                    workbook.create_sheet("View").add_chart(chart, "A1")
                elif hazard == "internal_location":
                    worksheet["B4"].hyperlink = Hyperlink(ref="B4", location="Records!A4")
                else:
                    worksheet["B4"].hyperlink = "#Records!A4"
                workbook.save(path)
                workbook.close()
                before = path.read_bytes()
                with self.assertRaises(import_config.ConfigFailure) as context:
                    import_config.sync_workbooks_from_schema(Path(directory), schema)
                self.assertIn("CONFIG_NOTE_MOVE_UNSAFE", {issue.code for issue in context.exception.issues})
                self.assertEqual(before, path.read_bytes())

    def test_generated_integer_input_aids_remain_safe_during_note_insertion(self):
        schema = note_schema(("Id", "Count", "Note"))
        with tempfile.TemporaryDirectory() as directory:
            workbook_dir = Path(directory)
            import_config.sync_workbooks_from_schema(workbook_dir, schema)
            path = workbook_dir / "ConfigTest.xlsx"
            workbook = import_config.load_workbook(path)
            worksheet = workbook["Records"]
            worksheet.append(["Keep", 4, "integer note"])
            workbook.save(path)
            workbook.close()
            updated = note_schema(("Id", "Count", "Extra", "Note"))
            self.assertTrue(import_config.sync_workbooks_from_schema(workbook_dir, updated))
            model = import_config.read_workbook_model(path, updated["layout"])
            rows, issues = import_config.validate_workbook_model(model, updated)
            self.assertEqual([], issues)
            self.assertEqual(4, rows["Records"][0]["Count"])
            self.assertEqual("integer note", rows["Records"][0]["Note"])


if __name__ == "__main__":
    unittest.main()
