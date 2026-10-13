from pathlib import Path
import hashlib
import json
from PIL import Image
from openpyxl import load_workbook

qa=Path(__file__).resolve().parent
repo=Path('F:/HunyueRoblox/TravelBuddy')
def read_tool(name):
    response=json.loads((qa/name).read_text(encoding='utf-8'))
    assert not response.get('isError'),name
    return json.loads(next(c['text'] for c in response['content'] if c['type']=='text'))
prepare=read_tool('preparing_verified.json')
paused=read_tool('paused_verified.json')
restored_pause=read_tool('paused_after_reload.json')
repack=read_tool('clear_repack_verified.json')
restored_prepare=read_tool('preparation_after_reload.json')
depart=read_tool('before_final_offline.json')
restored_journey=read_tool('final_journey_after_reload.json')
assert prepare['state']=='Preparing' and prepare['plan']['ReadyAt']>prepare['now']
assert paused['paused'] and 'ReadyAt' not in paused['plan']
assert paused['plan']['Id']==prepare['plan']['Id']==restored_pause['plan']['Id']
assert restored_pause['paused'] and restored_pause['now']>prepare['plan']['ReadyAt']
assert restored_pause['packed']==prepare['packed']==1
assert repack['empty']['packed']==0 and 'ReadyAt' not in repack['empty']['plan']
assert repack['repacked']['plan']['Id']==repack['before']['plan']['Id']==repack['empty']['plan']['Id']
assert repack['repacked']['plan']['SeaId']==repack['before']['plan']['SeaId']
assert repack['repacked']['plan']['DurationId']==repack['before']['plan']['DurationId']
assert repack['repacked']['plan']['ReadyAt']!=repack['before']['plan']['ReadyAt']
assert restored_prepare['state']=='Preparing' and restored_prepare['plan']==repack['repacked']['plan']
assert restored_prepare['warehouse']==2 and restored_prepare['packed']==1
assert depart['state']=='Voyaging' and depart['journey']['StartAt']==restored_prepare['plan']['ReadyAt']
assert depart['warehouse']==2 and depart['packed']==0
assert depart['journey']==restored_journey['journey']
assert restored_journey['warehouse']==2 and restored_journey['packed']==0 and restored_journey['oldDispatchDisabled']

before=json.loads((qa/'before.json').read_text(encoding='utf-8-sig'))
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()
framework={p:h for p,h in before.items() if '\\Framework\\' in p or p.endswith('StorageConfig.lua')}
assert all(sha(p)==h for p,h in framework.items())
changed={p:sha(p) for p,h in before.items() if sha(p)!=h}
log=(qa/'final-travel-demo.log').read_text(encoding='utf-8-sig')
assert 'TravelDomain: 31 passed, 0 failed' in log and 'TRAVEL_DEMO_CHECK_OK' in log
assert '0 errors' in (qa/'final-selene.log').read_text(encoding='utf-8-sig')
assert not (qa/'final-stylua.log').read_text(encoding='utf-8-sig').strip()

book=load_workbook(repo/'design/config/workbooks/TravelConfig.xlsx',data_only=False)
sheet=book['TravelSettings'];settings={sheet.cell(r,1).value:sheet.cell(r,2).value for r in range(4,sheet.max_row+1)}
for key,value in {'PrepareMinSeconds':60,'PrepareMaxSeconds':180,'FirstPrepareSeconds':30,'StudioPrepareMinSeconds':15,'StudioPrepareMaxSeconds':30,'RestSeconds':180,'SeaChoiceWeight_shallow':1,'SeaChoiceWeight_near':1,'ShortChoiceWeight':1,'LongChoiceWeight':1}.items():assert settings[key]==value,(key,settings[key])
image=qa.parents[1]/'assets/伙伴自主准备_Studio验证.png'
with Image.open(image) as png:png.verify()
report={'stage':'Saved preparation and autonomous departure','version':3,'framework_files_unchanged':len(framework),'changed_files':changed,'confirmed_settings':settings,'choice_weights':'Temporary adjustable Demo defaults; equal weights, not a frozen production decision','studio_evidence':{'prepare':prepare,'pause':paused,'paused_reload':restored_pause,'clear_and_repack':repack,'preparation_reload':restored_prepare,'departure':depart,'voyaging_reload':restored_journey},'checks':['31 Domain cases','64 ConfigTool cases','Service/World/Native/UI','ModuleBase/RojoWorkflow','StyLua/Selene','64 scripts compile','real Studio UI and ProfileStore reload'],'scope':'Wishes, automatic reward deposit and offline multiple voyages remain next stages; current ClaimReturn compatibility retained','test_asset_changes':'Two autonomous voyages consumed one packed supply each. Claimed existing returned voyages and opened one own sealed container through normal requests to test the next preparation. No profile reset, test currency injection or direct DataStore writes.','studio_final':'Play, own preparation table, Bag panel open; no temporary QA scripts retained','image':str(image)}
(qa/'verification-final.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print('AUTONOMY_FINAL_OK: saved plan, pause beyond old deadline, clear/repack same route, preparing restart equality, StartAt=ReadyAt, once-only supply debit, voyage restart equality, framework unchanged')
